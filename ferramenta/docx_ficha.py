"""Gera o .docx editavel de uma ficha a partir do modelo da coordenacao (data/docx_modelo/ficha_estudos_avalia_sesi_sp.docx).

Componente que falta: os conteudos aparecem com um '?' amarelo no lugar do icone e as contagens ficam com '?'.
Para concluir no Word: troque cada '?' pelo icone (copie da legenda) e acerte contagens/barra da pagina 1.
(O caminho mais facil e preencher alunos_editavel.xlsx e gerar de novo: o PDF sai identico ao oficial.)
"""
import copy, json, os
import docx
from docx.shared import Emu, Pt
from docx.text.paragraph import Paragraph
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls

TOOL = os.path.dirname(os.path.abspath(__file__))
MOD = os.path.join(TOOL, 'data', 'docx_modelo')
ICON = {'✓': 'image3.png', '~': 'image2.png', '!': 'image1.png'}
COR = {'✓': '1F8A5F', '~': 'C97F06', '!': 'CE4B39'}
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

def esc(s): return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

def set_text(p, text, keep_run=0):
    runs = p.findall(qn('w:r'))
    keep = runs[keep_run]
    for r in runs:
        if r is not keep: p.remove(r)
    for t in keep.findall(qn('w:t')) + keep.findall(qn('w:tab')): keep.remove(t)
    t = parse_xml(f'<w:t {W} xml:space="preserve">{esc(text)}</w:t>'); keep.append(t)

def fill_bullets(cell, items, shrink=False):
    ps = cell.findall(qn('w:p'))
    proto = copy.deepcopy(ps[0])
    if shrink:
        for e in proto.iter(qn('w:sz'), qn('w:szCs')): e.set(qn('w:val'), '12')
        proto.find(qn('w:pPr') + '/' + qn('w:spacing')).set(qn('w:line'), '158')
    for p in ps: cell.remove(p)
    for it in items:
        p = copy.deepcopy(proto); set_text(p, it); cell.append(p)

def run_xml(txt, color=None, sz=12, mono=False, hl=False):
    f = '<w:rFonts w:ascii="PlexMonoEstudos" w:eastAsia="PlexMonoEstudos" w:hAnsi="PlexMonoEstudos" w:cs="PlexMonoEstudos"/>' if mono else ''
    c = f'<w:color w:val="{color}"/>' if color else ''
    h = '<w:highlight w:val="yellow"/>' if hl else ''
    return f'<w:r><w:rPr>{f}{c}<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/>{h}</w:rPr><w:t xml:space="preserve">{esc(txt)}</w:t></w:r>'

def contagem_p(cnt, tem):
    g, a, r, tot = cnt
    out = f'<w:p {W}><w:pPr><w:spacing w:line="180" w:lineRule="exact"/></w:pPr>'
    for v, col in ((g, '1F8A5F'), (a, 'C97F06'), (r, 'CE4B39')):
        out += run_xml('●', col, 12) + run_xml(f' {v if tem else "?"}', '4A6470' if tem else None, 12, True, hl=not tem) + run_xml('    ', None, 12)
    out += run_xml(f'·  {tot} conteúdos', '4A6470', 12, True) + '</w:p>'
    return parse_xml(out)

def barra(cnt, tem, largura=6369):
    g, a, r, tot = cnt
    segs = [(v, c) for v, c in ((g, '1F8A5F'), (a, 'C97F06'), (r, 'CE4B39')) if v] if tem and tot else [(1, 'EEF1EF')]
    soma = sum(v for v, _ in segs); ws = [round(largura * v / soma) for v, _ in segs]; ws[-1] += largura - sum(ws)
    nob = ''.join(f'<w:{s} w:val="none" w:sz="0" w:space="0" w:color="FFFFFF"/>' for s in ('top', 'left', 'bottom', 'right'))
    cells = ''.join(f'<w:tc><w:tcPr><w:tcW w:w="{w_}" w:type="dxa"/><w:tcBorders>{nob}</w:tcBorders><w:shd w:val="clear" w:color="auto" w:fill="{c}"/>'
                    f'<w:tcMar><w:top w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tcMar></w:tcPr>'
                    f'<w:p><w:pPr><w:spacing w:after="0" w:line="20" w:lineRule="exact"/></w:pPr><w:r><w:rPr><w:sz w:val="2"/></w:rPr><w:t xml:space="preserve"> </w:t></w:r></w:p></w:tc>'
                    for (v, c), w_ in zip(segs, ws))
    grid = ''.join(f'<w:gridCol w:w="{w_}"/>' for w_ in ws)
    tbl = (f'<w:tbl {W}><w:tblPr><w:tblW w:w="{largura}" w:type="dxa"/><w:tblLayout w:type="fixed"/>'
           f'<w:tblCellMar><w:left w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tblCellMar></w:tblPr><w:tblGrid>{grid}</w:tblGrid>'
           f'<w:tr><w:trPr><w:trHeight w:hRule="exact" w:val="120"/></w:trPr>{cells}</w:tr></w:tbl>')
    return parse_xml(tbl)

def cartao(card, cnt, tem):
    cell = card.findall('.//' + qn('w:tc'))[0]
    ps = cell.findall(qn('w:p'))
    bar_p, cnt_p = ps[1], ps[2]
    idx = list(cell).index(bar_p)
    cell.remove(bar_p); cell.remove(cnt_p)
    cell.insert(idx, barra(cnt, tem))
    spacer = parse_xml(f'<w:p {W}><w:pPr><w:spacing w:after="0" w:line="100" w:lineRule="exact"/></w:pPr></w:p>')
    cell.insert(idx + 1, spacer); cell.insert(idx + 2, contagem_p(cnt, tem))

def direct(cell, tag): return [e for e in cell if e.tag == qn(tag)]

def metade(cell, doc, f, comp, D, tem):
    itens = f['itens'][comp]; desc = D['descritores'][comp]; n = len(desc)
    pars = direct(cell, 'w:p')
    cnt_p = next(p for p in pars if 'conteúdos avaliados' in ''.join(p.itertext()))
    set_text(cnt_p, f'{n} conteúdos avaliados' if tem else f'{n} conteúdos a preencher', keep_run=-1)
    tbls = direct(cell, 'w:tbl'); rows_t, tips_t = tbls[0], tbls[1]
    trs = rows_t.findall(qn('w:tr'))
    while len(trs) < n: nt = copy.deepcopy(trs[-1]); trs[-1].addnext(nt); trs = rows_t.findall(qn('w:tr'))
    for tr in trs[n:]: rows_t.remove(tr)
    trs = rows_t.findall(qn('w:tr'))
    h = 377 if n <= 10 else int(3770 / n)
    for i, tr in enumerate(trs):
        tr.find('.//' + qn('w:trHeight')).set(qn('w:val'), str(h))
        c0, c1 = tr.findall(qn('w:tc'))
        p0, p1 = c0.find(qn('w:p')), c1.find(qn('w:p'))
        for p in (p0, p1):
            for r in p.findall(qn('w:r')): p.remove(r)
        P0 = Paragraph(p0, doc._body)
        if tem:
            P0.add_run().add_picture(os.path.join(MOD, ICON[itens[i][0]]), width=Emu(142875), height=Emu(142875))
        else:
            p0.append(parse_xml(f'<w:r {W}><w:rPr><w:rFonts w:ascii="InterEstudos" w:hAnsi="InterEstudos"/><w:b/><w:sz w:val="15"/><w:highlight w:val="yellow"/></w:rPr><w:t xml:space="preserve"> ? </w:t></w:r>'))
        p1.append(parse_xml(f'<w:r {W}><w:rPr><w:rFonts w:ascii="InterEstudos" w:eastAsia="InterEstudos" w:hAnsi="InterEstudos" w:cs="InterEstudos"/><w:color w:val="4A6370"/><w:sz w:val="{15 if n <= 10 else 14}"/><w:szCs w:val="{15 if n <= 10 else 14}"/></w:rPr><w:t xml:space="preserve">{esc(desc[i])}</w:t></w:r>'))
    # dicas
    dic = D['dicas'][comp]
    title = next(p for p in tips_t.iter(qn('w:p')) if 'Para avan' in ''.join(p.itertext()) or 'Estrat' in ''.join(p.itertext()))
    set_text(title, dic['titulo'])
    cells = [c for c in tips_t.iter(qn('w:tc')) if c.find(qn('w:p') + '/' + qn('w:pPr') + '/' + qn('w:numPr')) is not None]
    big = max(sum(len(x) for x in col) for col in dic["colunas"]) > 700
    for c, items in zip(cells[:2], dic['colunas']): fill_bullets(c, items, big)

def gerar(f, path):
    D = json.load(open(os.path.join(TOOL, 'data', 'dados_por_ano.json')))[f['ano']]
    d = docx.Document(os.path.join(MOD, 'ficha_estudos_avalia_sesi_sp.docx'))
    t0, t1 = d.tables[0]._tbl, d.tables[1]._tbl
    R = t0.xpath('./w:tr/w:tc')[2]
    ch = list(R)
    spacer, greet, lp_card, mat_card, rot = ch[4], ch[5], ch[6], ch[8], ch[12]
    nome_p = parse_xml(f'<w:p {W}><w:pPr><w:spacing w:before="60" w:after="0" w:line="300" w:lineRule="exact"/></w:pPr>'
                       f'<w:r><w:rPr><w:rFonts w:ascii="BricolageEstudos" w:eastAsia="BricolageEstudos" w:hAnsi="BricolageEstudos" w:cs="BricolageEstudos"/><w:color w:val="17303A"/><w:sz w:val="25"/><w:szCs w:val="25"/></w:rPr><w:t xml:space="preserve">{esc(f["nome"])}</w:t></w:r></w:p>')
    ce_p = parse_xml(f'<w:p {W}><w:pPr><w:spacing w:after="0" w:line="220" w:lineRule="exact"/></w:pPr>' +
                     run_xml(f'CE 026 - Jaú · {f["ano"]} · Turma {f["turma"]} · RM {f["rm"]}', '4A6470', 13, True) + '</w:p>')
    spacer.addprevious(nome_p); spacer.addprevious(ce_p)
    spacer.find(qn('w:pPr') + '/' + qn('w:spacing')).set(qn('w:line'), '260')
    set_text(greet, f'Oi, {f["nome"].split()[0]}! Este mapa mostra como foi seu desempenho nos conteúdos avaliados e dá dicas de como continuar avançando nos estudos.')
    cartao(lp_card, f['contagem'][0], bool(f['itens']['LP'])); cartao(mat_card, f['contagem'][1], bool(f['itens']['MAT']))
    cells = [c for c in rot.iter(qn('w:tc'))]
    for c, items in zip(cells[1:3], D['rotina']): fill_bullets(c, items)
    # pagina 2
    tcs = t1.xpath('./w:tr')[0].findall(qn('w:tc'))
    hdr = next(p for p in direct(tcs[0], 'w:p') if p.xpath('string(.)').startswith('AVALIA'))
    set_text(hdr, f'AVALIA SESI-SP · {f["ano"]} {f["turma"]} · CE 026 - JAÚ')
    metade(tcs[0], d, f, 'LP', D, bool(f['itens']['LP']))
    metade(tcs[2], d, f, 'MAT', D, bool(f['itens']['MAT']))
    foot = t1.xpath('./w:tr')[1].findall(qn('w:tc'))[2].find(qn('w:p'))
    runs = foot.findall(qn('w:r')); last = runs[-1]
    for t in last.findall(qn('w:t')): last.remove(t)
    last.append(parse_xml(f'<w:t {W}>{esc(f["nome"])}</w:t>'))
    d.core_properties.title = f'Ficha de Estudos - {f["nome"]}'
    d.save(path)
