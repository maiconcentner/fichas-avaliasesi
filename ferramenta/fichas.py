#!/usr/bin/env python3
"""Gera fichas de estudos AvaliaSESI-SP (CE 026 Jau) identicas as da coordenacao.

Uso:
  python3 fichas.py gerar  ENTRADA.(json|xlsx) [-o PASTA] [--individual] [--docx]
  python3 fichas.py modelos [-o PASTA]
  python3 fichas.py conferir ENTRADA            (valida sem gerar)

ENTRADA = lista de alunos. Cada aluno:
  nome, ano ("9º EF", "1º EM"...), turma ("A"), rm,
  LP  = lista de situacoes na ordem oficial dos conteudos do ano (✓ ~ !)  ou vazia/ausente se falta o componente
  MAT = idem
A situacao tambem pode vir como [acertos, total] (todas certas = ✓, nenhuma = !, demais = ~).
O xlsx gerado (alunos_editavel.xlsx) e o proprio formato de entrada: preencha o que falta e rode `gerar` de novo.
"""
import argparse, json, os, re, sys, unicodedata
TOOL = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, TOOL)

SIM = ('✓', '~', '!')
NOME_SIM = {'✓': 'Você domina', '~': 'Quase lá', '!': 'Comece por aqui'}
ORDEM_ANO = ['3º EF', '4º EF', '5º EF', '6º EF', '7º EF', '8º EF', '9º EF', '1º EM', '2º EM', '3º EM']

def dados():
    return json.load(open(os.path.join(TOOL, 'data', 'dados_por_ano.json')))

def norm_ano(v):
    s = unicodedata.normalize('NFKD', str(v)).encode('ascii', 'ignore').decode().lower()
    m = re.search(r'(\d)', s)
    if not m: raise ValueError(f'ano invalido: {v!r}')
    n = int(m.group(1))
    em = bool(re.search(r'\bem\b|medio|serie', s))
    ef = bool(re.search(r'\bef\b|fundamental', s))
    seg = 'EM' if em else 'EF' if ef or n >= 4 else None
    if seg is None: raise ValueError(f'ano ambiguo (informe EF ou EM): {v!r}')
    ano = f'{n}º {seg}'
    if ano not in ORDEM_ANO: raise ValueError(f'ano sem modelo: {ano}')
    return ano

def sit(v):
    """normaliza uma situacao: simbolo, texto ou [acertos, total]"""
    if v is None: return None
    if isinstance(v, (list, tuple)):
        a, t = v
        if t in (None, 0, ''): return None
        return '✓' if a >= t else '!' if a <= 0 else '~'
    s = str(v).strip()
    if s == '': return None
    if s in SIM: return s
    k = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower()
    if k.startswith(('voce domina', 'domina', 'dominio', 'ok', 'certo')) or s in ('✔', '√'): return '✓'
    if k.startswith(('quase', 'parcial')) or s == '≈': return '~'
    if k.startswith(('comece', 'nao domina', 'erro', 'errado')) or s in ('❗', '¡'): return '!'
    m = re.fullmatch(r'(\d+)\s*/\s*(\d+)', s)
    if m: return sit([int(m.group(1)), int(m.group(2))])
    raise ValueError(f'situacao invalida: {v!r}')

def contagem(itens):
    return [[sum(1 for s, _ in itens[k] if s == sy) for sy in SIM] + [len(itens[k])] for k in ('LP', 'MAT')]

def normalizar(rec, D=None):
    """aluno bruto -> ficha pronta para gen.make_pdf. Levanta ValueError com mensagem clara."""
    D = D or dados()
    nome = ' '.join(str(rec['nome']).split())
    ano = norm_ano(rec['ano'])
    turma = str(rec['turma']).strip().upper()
    rm = re.sub(r'\.0$', '', str(rec['rm']).strip())
    if not nome or not turma or not rm.isdigit(): raise ValueError(f'{nome or "?"}: nome, turma e RM (numerico) sao obrigatorios')
    itens = {}
    for k in ('LP', 'MAT'):
        desc = D[ano]['descritores'][k]
        st = rec.get(k) or []
        st = [sit(x) for x in st]
        if all(x is None for x in st): itens[k] = []; continue
        if len(st) != len(desc): raise ValueError(f'{nome} ({ano}) {k}: {len(st)} situacoes, o ano tem {len(desc)} conteudos')
        if any(x is None for x in st): raise ValueError(f'{nome} ({ano}) {k}: ha conteudos sem situacao (preencha todos ou deixe o componente vazio)')
        itens[k] = [[s, d] for s, d in zip(st, desc)]
    f = dict(nome=nome, ano=ano, turma=turma, rm=rm, itens=itens, contagem=contagem(itens))
    f['situacao'] = 'completo' if itens['LP'] and itens['MAT'] else 'falta LP' if itens['MAT'] else 'falta MAT' if itens['LP'] else 'sem dados'
    return f

# ---------- leitura/escrita ----------
def ler(path):
    if path.lower().endswith('.json'):
        raw = json.load(open(path, encoding='utf8'))
        return raw['alunos'] if isinstance(raw, dict) else raw
    return ler_xlsx(path)

def ler_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Respostas'] if 'Respostas' in wb.sheetnames else wb.worksheets[0]
    hdr = [str(c.value or '').strip().lower() for c in ws[1]]
    ix = {h: i for i, h in enumerate(hdr)}
    need = ['nome', 'ano', 'turma', 'rm', 'componente', 'situação']
    miss = [h for h in need if h not in ix]
    if miss: raise ValueError(f'colunas ausentes na planilha: {miss}')
    alunos = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[ix['nome']]: continue
        key = (str(row[ix['rm']]), str(row[ix['nome']]))
        a = alunos.setdefault(key, dict(nome=row[ix['nome']], ano=row[ix['ano']], turma=row[ix['turma']], rm=row[ix['rm']], LP=[], MAT=[]))
        comp = str(row[ix['componente']]).strip().upper()
        comp = 'LP' if comp.startswith('L') else 'MAT'
        a[comp].append(row[ix['situação']])
    return list(alunos.values())

def escrever_xlsx(fichas, path):
    """planilha editavel: uma linha por conteudo; situacao vazia = a preencher"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule
    D = dados()
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Respostas'
    ws.append(['Nome', 'Ano', 'Turma', 'RM', 'Componente', 'Nº', 'Conteúdo', 'Situação'])
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='17303A')
    for f in fichas:
        for k, nomek in (('LP', 'LP'), ('MAT', 'MAT')):
            desc = D[f['ano']]['descritores'][k]
            it = f['itens'][k]
            for n, d in enumerate(desc, 1):
                ws.append([f['nome'], f['ano'], f['turma'], int(f['rm']), nomek, n, d, it[n - 1][0] if it else None])
    n = ws.max_row
    dv = DataValidation(type='list', formula1='"✓,~,!"', allow_blank=True, showErrorMessage=True, errorTitle='Situação', error='Use ✓ (domina), ~ (quase lá) ou ! (comece por aqui).')
    ws.add_data_validation(dv); dv.add(f'H2:H{n}')
    ws.conditional_formatting.add(f'H2:H{n}', CellIsRule(operator='equal', formula=['"✓"'], fill=PatternFill('solid', bgColor='C8E6D8')))
    ws.conditional_formatting.add(f'H2:H{n}', CellIsRule(operator='equal', formula=['"~"'], fill=PatternFill('solid', bgColor='F6DFB0')))
    ws.conditional_formatting.add(f'H2:H{n}', CellIsRule(operator='equal', formula=['"!"'], fill=PatternFill('solid', bgColor='F3C6BF')))
    ws.conditional_formatting.add(f'H2:H{n}', CellIsRule(operator='equal', formula=['""'], fill=PatternFill('solid', bgColor='FFF2A8')))
    for col, w in zip('ABCDEFGH', (34, 8, 7, 8, 12, 5, 80, 10)): ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=2): r[7].alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:H{n}'
    r2 = wb.create_sheet('Alunos')
    r2.append(['Nome', 'Ano', 'Turma', 'RM', 'Situação da ficha', 'LP ✓ ~ !', 'MAT ✓ ~ !'])
    for c in r2[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='17303A')
    for f in fichas:
        c = f['contagem']
        r2.append([f['nome'], f['ano'], f['turma'], int(f['rm']), f['situacao'], ' '.join(map(str, c[0][:3])) if f['itens']['LP'] else '—', ' '.join(map(str, c[1][:3])) if f['itens']['MAT'] else '—'])
    for col, w in zip('ABCDEFG', (34, 8, 7, 8, 18, 12, 12)): r2.column_dimensions[col].width = w
    r3 = wb.create_sheet('Leia-me')
    for line in ['Como usar', '1) Na aba Respostas, preencha a coluna Situação das linhas amarelas (componente que falta).',
                 '2) Use ✓ (Você domina), ~ (Quase lá) ou ! (Comece por aqui). Regra: acertou todas = ✓; acertou parte = ~; nenhuma = !.',
                 '3) Não altere Nome, Ano, Turma, RM nem a ordem dos conteúdos.',
                 '4) Salve e gere novamente: python3 fichas.py gerar alunos_editavel.xlsx']: r3.append([line])
    r3.column_dimensions['A'].width = 120; r3['A1'].font = Font(bold=True, size=13)
    wb.save(path)

# ---------- comandos ----------
def preparar(path):
    D = dados(); out = []; erros = []
    for rec in ler(path):
        try: out.append(normalizar(rec, D))
        except (ValueError, KeyError) as e: erros.append(str(e))
    out.sort(key=lambda f: (ORDEM_ANO.index(f['ano']), f['turma'], f['nome']))
    try:
        of = {s['rm']: s for s in json.load(open(os.path.join(TOOL, 'data', 'alunos_oficiais.json')))}
        for f in out:
            if f['rm'] in of: print(f"AVISO: {f['nome']} (RM {f['rm']}) ja consta no relatorio oficial ({of[f['rm']]['arquivo']}, componentes: LP={'sim' if of[f['rm']]['itens']['LP'] else 'nao'} MAT={'sim' if of[f['rm']]['itens']['MAT'] else 'nao'})")
    except FileNotFoundError: pass
    return out, erros

def slug(s): return re.sub(r'[^A-Za-z0-9]+', '_', unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode()).strip('_')

def cmd_gerar(a):
    import gen
    fichas, erros = preparar(a.entrada)
    for e in erros: print('ERRO:', e, file=sys.stderr)
    if erros and not a.ignorar_erros: sys.exit(2)
    os.makedirs(a.o, exist_ok=True)
    sem = [f for f in fichas if f['situacao'] == 'sem dados']
    for f in sem: print('AVISO: sem dados de LP nem MAT:', f['nome'])
    fichas = [f for f in fichas if f['situacao'] != 'sem dados']
    gen.make_pdf(fichas, os.path.join(a.o, 'fichas_geradas.pdf'), os.path.join(a.o, '_fichas.html'))
    os.remove(os.path.join(a.o, '_fichas.html'))
    if a.individual:
        d = os.path.join(a.o, 'individuais'); os.makedirs(d, exist_ok=True)
        for f in fichas:
            gen.make_pdf([f], os.path.join(d, f"{f['ano'].replace('º ', '')}_{f['turma']}_{slug(f['nome'])}.pdf"), os.path.join(d, '_t.html'))
        if os.path.exists(os.path.join(d, '_t.html')): os.remove(os.path.join(d, '_t.html'))
    parc = [f for f in fichas if f['situacao'] != 'completo']
    escrever_xlsx(fichas, os.path.join(a.o, 'alunos_editavel.xlsx'))
    if a.docx:
        import docx_ficha
        d = os.path.join(a.o, 'editaveis_docx'); os.makedirs(d, exist_ok=True)
        for f in (fichas if a.docx == 'todos' else parc):
            docx_ficha.gerar(f, os.path.join(d, f"{f['ano'].replace('º ', '')}_{f['turma']}_{slug(f['nome'])}.docx"))
    print(f'{len(fichas)} fichas ({len(fichas) - len(parc)} completas, {len(parc)} parciais) em {a.o}')
    for f in parc: print('  parcial:', f['nome'], f['ano'], f['turma'], '->', f['situacao'])

def cmd_conferir(a):
    fichas, erros = preparar(a.entrada)
    for e in erros: print('ERRO:', e)
    for f in fichas: print(f['situacao'].ljust(10), f['ano'], f['turma'], f['nome'], f['contagem'])
    sys.exit(1 if erros else 0)

def exemplo(ano, D):
    ciclo = ['✓', '~', '!']
    return dict(nome='Nome do Estudante', ano=ano, turma='A', rm='0000',
                LP=[ciclo[i % 3] for i in range(len(D[ano]['descritores']['LP']))],
                MAT=[ciclo[(i + 1) % 3] for i in range(len(D[ano]['descritores']['MAT']))])

def cmd_modelos(a):
    import gen
    D = dados(); os.makedirs(a.o, exist_ok=True)
    for ano in ORDEM_ANO:
        f = normalizar(exemplo(ano, D), D)
        gen.make_pdf([f], os.path.join(a.o, f"modelo_{ano.replace('º ', '')}.pdf"), os.path.join(a.o, '_m.html'))
    os.remove(os.path.join(a.o, '_m.html'))
    escrever_xlsx([normalizar({**exemplo(ano, D), 'nome': f'Exemplo {ano}'}, D) for ano in ORDEM_ANO], os.path.join(a.o, 'modelo_planilha_entrada.xlsx'))
    import docx_ficha
    for ano in ORDEM_ANO:
        docx_ficha.gerar(normalizar(exemplo(ano, D), D), os.path.join(a.o, f"modelo_{ano.replace('º ', '')}.docx"))
    print('modelos gerados em', a.o)

if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = p.add_subparsers(dest='cmd', required=True)
    g = sp.add_parser('gerar'); g.add_argument('entrada'); g.add_argument('-o', default='saida_fichas'); g.add_argument('--individual', action='store_true')
    g.add_argument('--docx', nargs='?', const='parciais', choices=['parciais', 'todos']); g.add_argument('--ignorar-erros', action='store_true'); g.set_defaults(f=cmd_gerar)
    c = sp.add_parser('conferir'); c.add_argument('entrada'); c.set_defaults(f=cmd_conferir)
    m = sp.add_parser('modelos'); m.add_argument('-o', default='modelos_fichas'); m.set_defaults(f=cmd_modelos)
    a = p.parse_args(); a.f(a)
